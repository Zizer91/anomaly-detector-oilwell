"""
Streamlit-дашборд для детекции аномалий в данных нефтяных скважин.
Загружает parquet-файл, применяет обученную модель и визуализирует результаты.
"""

import streamlit as st
import pandas as pd
import numpy as np
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import joblib
from pathlib import Path
from io import StringIO
from src.data_preprocessor import DataPreprocessor

# ============================================================
# КОНФИГУРАЦИЯ (должна совпадать с обученной моделью)
# ============================================================
FEATURE_COLS = ['P-PDG', 'P-TPT', 'T-PDG', 'T-TPT', 'T-JUS-CKP', 'T-MON-CKP']
WINDOW_SIZE = 120
STEP = 30
CONTAMINATION = 0.1

# ============================================================
# ЗАГРУЗКА МОДЕЛИ
# ============================================================
@st.cache_resource
def load_detector():
    """Загружает обученную модель."""
    model_path = Path('anomaly_detector.pkl')
    if not model_path.exists():
        st.error("Модель не найдена! Сначала обучите модель через train.py")
        return None
    return joblib.load(model_path)

# ============================================================
# ПРЕДОБРАБОТКА ДАННЫХ
# ============================================================
def preprocess_data(df: pd.DataFrame) -> tuple:
    """
    Подготавливает данные для модели.
    Возвращает окна, предсказания и метки аномалий для каждого окна.
    """
    prep = DataPreprocessor(window_size=WINDOW_SIZE, step=STEP)

    # Проверяем наличие всех признаков
    available_features = [f for f in FEATURE_COLS if f in df.columns]
    missing_features = set(FEATURE_COLS) - set(available_features)
    if missing_features:
        st.warning(f"Отсутствуют признаки: {missing_features}")

    if len(available_features) < 2:
        st.error("Недостаточно признаков для анализа")
        return None, None, None

    # Нарезаем окна
    X = prep.create_windows(df, available_features)
    if len(X) == 0:
        st.error("Недостаточно данных для создания окон")
        return None, None, None

    # Загружаем модель и предсказываем
    detector = load_detector()
    if detector is None:
        return None, None, None

    preds = detector.predict(X.reshape(X.shape[0], -1))
    anomaly_labels = [1 if p == -1 else 0 for p in preds]

    return X, preds, anomaly_labels


def plot_results(df: pd.DataFrame, anomaly_labels: list, window_size: int, step: int):
    """Строит интерактивные графики результатов детекции."""

    # Создаем массив меток для каждой точки (для отображения)
    point_labels = np.zeros(len(df))
    for i, is_anomaly in enumerate(anomaly_labels):
        start_idx = i * step
        end_idx = min(start_idx + window_size, len(df))
        if is_anomaly == 1:
            point_labels[start_idx:end_idx] = 1

    # Основные параметры для отображения
    params_to_plot = ['P-PDG', 'P-TPT', 'T-PDG', 'QGL']
    available_params = [p for p in params_to_plot if p in df.columns]

    if not available_params:
        available_params = [df.columns[0]]  # первый доступный параметр

    fig = make_subplots(
        rows=len(available_params),
        cols=1,
        shared_xaxes=True,
        vertical_spacing=0.05,
        subplot_titles=available_params
    )

    colors = ['#1f77b4', '#ff7f0e', '#2ca02c', '#d62728', '#9467bd']

    for idx, param in enumerate(available_params, 1):
        # График параметра
        fig.add_trace(
            go.Scatter(
                x=df.index,
                y=df[param],
                mode='lines',
                name=param,
                line=dict(color=colors[idx % len(colors)], width=1.5),
                showlegend=False
            ),
            row=idx, col=1
        )

        # Выделение аномалий
        anomaly_indices = df.index[point_labels == 1]
        anomaly_values = df[param][point_labels == 1]

        if len(anomaly_indices) > 0:
            fig.add_trace(
                go.Scatter(
                    x=anomaly_indices,
                    y=anomaly_values,
                    mode='markers',
                    name='Аномалия' if idx == 1 else None,
                    marker=dict(color='red', size=4, symbol='circle'),
                    showlegend=(idx == 1)
                ),
                row=idx, col=1
            )

        # Настройки осей
        fig.update_yaxes(title_text=param, row=idx, col=1)

    fig.update_xaxes(title_text='Время', row=len(available_params), col=1)
    fig.update_layout(
        height=300 * len(available_params),
        title_text="Детекция аномалий",
        showlegend=True,
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1)
    )

    st.plotly_chart(fig, use_container_width=True)

    # Дополнительная статистика
    st.subheader("📊 Статистика детекции")

    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Всего окон", len(anomaly_labels))
    with col2:
        st.metric("Аномальных окон", sum(anomaly_labels))
    with col3:
        anomaly_rate = sum(anomaly_labels) / len(anomaly_labels) if anomaly_labels else 0
        st.metric("Доля аномалий", f"{anomaly_rate:.1%}")


def export_results(df: pd.DataFrame, anomaly_labels: list, window_size: int, step: int):
    """Экспортирует результаты детекции в CSV и Excel."""

    # Создаём детальную таблицу по окнам
    windows_data = []
    for i, is_anomaly in enumerate(anomaly_labels):
        start_idx = i * step
        end_idx = min(start_idx + window_size, len(df))

        window_info = {
            'window_id': i + 1,
            'start_time': df.index[start_idx] if len(df.index) > start_idx else None,
            'end_time': df.index[end_idx - 1] if end_idx > 0 and len(df.index) > end_idx - 1 else None,
            'is_anomaly': is_anomaly,
            'anomaly_label': 'Аномалия' if is_anomaly else 'Норма'
        }
        windows_data.append(window_info)

    results_df = pd.DataFrame(windows_data)

    # Создаём CSV
    csv_buffer = StringIO()
    results_df.to_csv(csv_buffer, index=False, encoding='utf-8-sig')
    csv_data = csv_buffer.getvalue()

    # Создаём Excel (если openpyxl установлен)
    excel_data = None
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Detection Results"

        # Заголовки
        headers = ['№ окна', 'Начало', 'Конец', 'Аномалия', 'Статус']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)

        # Данные
        for row, data in enumerate(windows_data, 2):
            ws.cell(row=row, column=1, value=data['window_id'])
            ws.cell(row=row, column=2, value=str(data['start_time']) if data['start_time'] else '')
            ws.cell(row=row, column=3, value=str(data['end_time']) if data['end_time'] else '')
            ws.cell(row=row, column=4, value=1 if data['is_anomaly'] else 0)
            ws.cell(row=row, column=5, value=data['anomaly_label'])

            # Подсветка аномалий
            if data['is_anomaly']:
                for col in range(1, 6):
                    ws.cell(row=row, column=col).fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")

        # Автоширина колонок
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[col_letter].width = adjusted_width

        from io import BytesIO
        excel_buffer = BytesIO()
        wb.save(excel_buffer)
        excel_data = excel_buffer.getvalue()

    except ImportError:
        st.info("Для экспорта в Excel установите: pip install openpyxl")

    # Кнопки экспорта
    col1, col2 = st.columns(2)

    with col1:
        st.download_button(
            label="📥 Скачать как CSV",
            data=csv_data,
            file_name="anomaly_detection_results.csv",
            mime="text/csv",
            key="csv_download"
        )

    with col2:
        if excel_data:
            st.download_button(
                label="📊 Скачать как Excel",
                data=excel_data,
                file_name="anomaly_detection_results.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="excel_download"
            )

    # Показываем预览 таблицы
    with st.expander("📋 Предпросмотр результатов"):
        st.dataframe(results_df.head(20), use_container_width=True)
        st.caption(f"Всего записей: {len(results_df)}")


# ============================================================
# ОСНОВНОЙ ИНТЕРФЕЙС
# ============================================================
st.set_page_config(page_title="Anomaly Detector", layout="wide")

st.title("🛢️ Детектор аномалий в работе нефтяных скважин")
st.markdown("---")

# Боковая панель
with st.sidebar:
    st.header("📁 Загрузка данных")
    uploaded_file = st.file_uploader(
        "Выберите файл в формате Parquet",
        type=['parquet'],
        help="Поддерживаются файлы из датасета 3W (Petrobras)"
    )

    st.header("⚙️ Параметры модели")
    st.info(f"""
    - Размер окна: {WINDOW_SIZE} сек
    - Шаг: {STEP} сек
    - Доля аномалий: {CONTAMINATION}
    - Признаки: {', '.join(FEATURE_COLS[:4])}...
    """)

    st.header("ℹ️ О программе")
    st.markdown("""
    Детектор аномалий на основе Isolation Forest.
    
    **Обучен на:**  
    - Класс 5 (быстрая потеря продуктивности)
    - F1-score: 0.985
    
    **Как использовать:**
    1. Загрузите parquet-файл
    2. Модель автоматически проанализирует данные
    3. Аномалии отмечены красными точками
    4. Скачайте результаты в CSV или Excel
    """)

# Основная область
if uploaded_file is not None:
    with st.spinner("Загрузка и анализ данных..."):
        try:
            # Чтение данных
            df = pd.read_parquet(uploaded_file)
            st.success(f"✅ Данные загружены: {len(df)} строк, {len(df.columns)} колонок")

            # Создание временного индекса (если его нет)
            if not isinstance(df.index, pd.DatetimeIndex):
                if 'timestamp' in df.columns:
                    df['timestamp'] = pd.to_datetime(df['timestamp'])
                    df.set_index('timestamp', inplace=True)
                else:
                    # Создаем искусственный индекс с шагом 1 секунда
                    df.index = pd.date_range(
                        start='2024-01-01 00:00:00',
                        periods=len(df),
                        freq='S'
                    )

            # Предобработка и предсказание
            X, preds, anomaly_labels = preprocess_data(df)

            if X is not None:
                # Визуализация
                plot_results(df, anomaly_labels, WINDOW_SIZE, STEP)

                # Экспорт результатов
                st.subheader("📎 Экспорт результатов")
                export_results(df, anomaly_labels, WINDOW_SIZE, STEP)

        except Exception as e:
            st.error(f"Ошибка при обработке файла: {str(e)}")
else:
    st.info("👈 Загрузите parquet-файл с данными скважины для анализа")

    # Пример данных (показываем, чего ожидать)
    with st.expander("📖 Пример формата данных"):
        st.markdown("""
        **Ожидаемые колонки:**
        - `P-PDG` - забойное давление
        - `P-TPT` - давление на TPT
        - `T-PDG` - температура на PDG
        - `T-TPT` - температура на TPT
        - `T-JUS-CKP` - температура после штуцера
        - `T-MON-CKP` - температура перед штуцером
        
        **Другие колонки** (опционально):
        - `QGL` - дебит газа
        - `class` - метка класса (если есть, для валидации)
        """)

st.markdown("---")
st.caption("Детектор аномалий | Обучен на данных Petrobras 3W | Быстрая потеря продуктивности (Class 5)")